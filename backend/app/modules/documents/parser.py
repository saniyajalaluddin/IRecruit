"""Document parsing and text normalization pipeline for PDF, DOCX, and TXT files."""

import io
import re
import unicodedata
from typing import Optional
import docx
import pypdf

from backend.app.core.errors import UnsupportedDocumentError
from backend.app.modules.documents.schemas import DocumentFormat, DocumentMetadata, ParsedDocument


class DocumentParser:
    """Extracts, validates, and normalizes machine-readable plaintext from document formats."""

    @classmethod
    def normalize_extracted_text(cls, text: str) -> str:
        """Normalizes unicode characters, standardizes bullets, and strips extraneous whitespace."""
        if not text:
            return ""

        # Normalize unicode to NFKC
        text = unicodedata.normalize("NFKC", text)

        # Standardize various unicode bullet points to standard bullet '- '
        text = re.sub(r"[\u2022\u2023\u25E6\u2043\u2219\u25CB\u25CF\u25AA\u25A0]\s*", "- ", text)

        # Replace non-breaking spaces and tabs
        text = text.replace("\xa0", " ").replace("\t", " ")

        # Remove null bytes and non-printable control characters (retain \n and \r)
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

        # Standardize carriage returns
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        # Collapse multiple spaces on the same line
        text = re.sub(r"[ ]{2,}", " ", text)

        # Collapse excess blank lines (more than 2 consecutive newlines)
        text = re.sub(r"\n{3,}", "\n\n", text)

        # Strip trailing/leading whitespace per line
        lines = [line.strip() for line in text.split("\n")]
        text = "\n".join(lines).strip()

        return text

    @classmethod
    def parse_pdf(cls, content: bytes) -> tuple[str, int]:
        """Extracts text from PDF, detecting scanned/image-only pages."""
        bio = io.BytesIO(content)
        try:
            reader = pypdf.PdfReader(bio)
        except Exception as e:
            raise UnsupportedDocumentError(f"Failed to open PDF document: {str(e)}")

        if reader.is_encrypted:
            raise UnsupportedDocumentError("Encrypted or password-protected PDF files cannot be processed.")

        page_count = len(reader.pages)
        if page_count == 0:
            raise UnsupportedDocumentError("PDF document contains no pages.")

        extracted_parts = []
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
                extracted_parts.append(page_text.strip())
            except Exception as e:
                # Page level extraction failure
                raise UnsupportedDocumentError(f"Failed extracting text on page {i + 1}: {str(e)}")

        combined_text = "\n\n".join([p for p in extracted_parts if p]).strip()

        # Image-only / Scanned PDF detection
        if len(combined_text) < 15:
            raise UnsupportedDocumentError(
                "This PDF appears to be a scanned or image-only document without a machine-readable text layer. "
                "Please submit a text-based PDF or DOCX."
            )

        return combined_text, page_count

    @classmethod
    def parse_docx(cls, content: bytes) -> tuple[str, int]:
        """Extracts text from DOCX paragraphs and tables."""
        bio = io.BytesIO(content)
        try:
            doc = docx.Document(bio)
        except Exception as e:
            raise UnsupportedDocumentError(f"Failed to open DOCX document: {str(e)}")

        parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())

        # Extract tables (e.g. skills or experience listed in table format)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate repeated cell text from merged cells
                    unique_cells = list(dict.fromkeys(row_cells))
                    parts.append(" | ".join(unique_cells))

        combined_text = "\n\n".join(parts).strip()
        if not combined_text:
            raise UnsupportedDocumentError("The DOCX document contains no readable text.")

        return combined_text, 1

    @classmethod
    def parse_txt(cls, content: bytes) -> tuple[str, int]:
        """Decodes and extracts text from TXT files."""
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception:
                raise UnsupportedDocumentError("Text document could not be decoded. Please ensure UTF-8 encoding.")

        cleaned = text.strip()
        if not cleaned:
            raise UnsupportedDocumentError("The text document is empty.")

        return cleaned, 1

    @classmethod
    def parse(
        cls,
        content: bytes,
        metadata: DocumentMetadata,
    ) -> ParsedDocument:
        """Parses binary content into normalized structured ParsedDocument."""
        if metadata.format == DocumentFormat.PDF:
            raw_text, page_count = cls.parse_pdf(content)
        elif metadata.format == DocumentFormat.DOCX:
            raw_text, page_count = cls.parse_docx(content)
        elif metadata.format == DocumentFormat.TXT:
            raw_text, page_count = cls.parse_txt(content)
        else:
            raise UnsupportedDocumentError(f"Unsupported format: {metadata.format}")

        normalized_text = cls.normalize_extracted_text(raw_text)
        word_count = len(normalized_text.split())

        return ParsedDocument(
            raw_text=raw_text,
            normalized_text=normalized_text,
            page_count=page_count,
            word_count=word_count,
            metadata=metadata,
            is_empty=word_count == 0,
        )
