"""Unit tests for document parsing, normalization, and failure detections."""

import io
import pytest
import pypdf
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import docx
from backend.app.core.errors import UnsupportedDocumentError
from backend.app.modules.documents.parser import DocumentParser
from backend.app.modules.documents.schemas import DocumentFormat, DocumentMetadata


def create_pdf_with_text(text: str) -> bytes:
    """Creates a valid PDF with real standard Type1 font stream text."""
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=300, height=300)

    # Attach standard Type1 font resource
    font_dict = DictionaryObject()
    font_dict[NameObject("/Type")] = NameObject("/Font")
    font_dict[NameObject("/Subtype")] = NameObject("/Type1")
    font_dict[NameObject("/BaseFont")] = NameObject("/Helvetica")

    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font_dict

    res = DictionaryObject()
    res[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = res

    # Stream content
    escaped = text.replace("(", "\\(").replace(")", "\\)")
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 50 250 Td ({escaped}) Tj ET".encode("utf-8"))
    page[NameObject("/Contents")] = stream

    bio = io.BytesIO()
    writer.write(bio)
    return bio.getvalue()


def create_blank_scanned_pdf() -> bytes:
    """Creates a blank/image-only PDF with zero text."""
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=300, height=300)
    bio = io.BytesIO()
    writer.write(bio)
    return bio.getvalue()


def create_docx_with_tables() -> bytes:
    """Creates a DOCX with paragraphs and a skill table."""
    doc = docx.Document()
    doc.add_heading("Jane Doe - Resume", level=1)
    doc.add_paragraph("Senior Python Backend Engineer with 6 years experience.")
    
    # Add table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Technologies"
    table.cell(1, 0).text = "Languages"
    table.cell(1, 1).text = "Python, TypeScript, SQL"

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_text_normalization():
    """Verify unicode NFKC normalization, bullet point conversion, and whitespace collapse."""
    raw = "• Python\tDeveloper\r\n\xa0\xa0\u2022 FastAPI & Docker\n\n\n\n- PostgreSQL"
    normalized = DocumentParser.normalize_extracted_text(raw)
    assert "•" not in normalized
    assert "- Python Developer" in normalized
    assert "- FastAPI & Docker" in normalized
    assert "- PostgreSQL" in normalized
    assert "\t" not in normalized
    assert "\xa0" not in normalized
    # Max two consecutive newlines
    assert "\n\n\n" not in normalized


def test_docx_parsing_with_tables():
    """Verify DOCX parsing captures both paragraph text and table cell content."""
    docx_bytes = create_docx_with_tables()
    metadata = DocumentMetadata(
        original_filename="jane_resume.docx",
        sanitized_filename="jane_resume.docx",
        format=DocumentFormat.DOCX,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size_bytes=len(docx_bytes),
    )

    parsed = DocumentParser.parse(docx_bytes, metadata)
    assert parsed.is_empty is False
    assert "Jane Doe" in parsed.normalized_text
    assert "Senior Python Backend Engineer" in parsed.normalized_text
    assert "Python, TypeScript, SQL" in parsed.normalized_text
    assert parsed.word_count > 10


def test_txt_parsing():
    """Verify TXT file decoding and word count."""
    text_content = "Candidate Alice Cooper\nExperience: 5 years in Distributed Systems\nSkills: Python, Go, Kafka"
    metadata = DocumentMetadata(
        original_filename="alice.txt",
        sanitized_filename="alice.txt",
        format=DocumentFormat.TXT,
        content_type="text/plain",
        file_size_bytes=len(text_content.encode("utf-8")),
    )
    parsed = DocumentParser.parse(text_content.encode("utf-8"), metadata)
    assert "Distributed Systems" in parsed.normalized_text
    assert parsed.word_count == 13


def test_image_only_pdf_detection():
    """Verify image-only / scanned PDF without text layer raises UnsupportedDocumentError."""
    blank_pdf_bytes = create_blank_scanned_pdf()
    metadata = DocumentMetadata(
        original_filename="scanned_resume.pdf",
        sanitized_filename="scanned_resume.pdf",
        format=DocumentFormat.PDF,
        content_type="application/pdf",
        file_size_bytes=len(blank_pdf_bytes),
    )

    with pytest.raises(UnsupportedDocumentError) as excinfo:
        DocumentParser.parse(blank_pdf_bytes, metadata)
    assert "scanned or image-only" in str(excinfo.value)


def test_empty_txt_file_rejected():
    """Verify empty text file raises UnsupportedDocumentError."""
    metadata = DocumentMetadata(
        original_filename="empty.txt",
        sanitized_filename="empty.txt",
        format=DocumentFormat.TXT,
        content_type="text/plain",
        file_size_bytes=0,
    )
    with pytest.raises(UnsupportedDocumentError):
        DocumentParser.parse(b"   \n\t  ", metadata)


def test_pdf_with_text_extraction():
    """Verify parsing a PDF containing text extracts and normalizes properly."""
    pdf_bytes = create_pdf_with_text("Senior Software Engineer with Python experience")
    metadata = DocumentMetadata(
        original_filename="engineer.pdf",
        sanitized_filename="engineer.pdf",
        format=DocumentFormat.PDF,
        content_type="application/pdf",
        file_size_bytes=len(pdf_bytes),
    )
    parsed = DocumentParser.parse(pdf_bytes, metadata)
    assert "Senior Software Engineer" in parsed.normalized_text
    assert parsed.page_count == 1
    assert parsed.word_count == 6
